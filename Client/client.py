import sys
import socket
from PyQt5 import QtWidgets,uic
from PyQt5.QtWidgets import QDialog, QApplication,QPushButton
from PyQt5.uic import loadUi
from fpdf import FPDF
import docx
import os
import win32com.client

class CWordAutomate:
    def __init__( self ):
        self.obWord         = win32com.client.Dispatch( "Word.Application" )
        self.obDoc          = self.obWord.Documents.Add( )
        self.obWord.Visible = 1
        self.Sel            = self.obWord.Selection

    def Write( self, text, font, size, bold=0 ):
        self.Sel.Font.Name = font
        self.Sel.Font.Bold = bold
        self.Sel.Font.Size = size
        self.Sel.TypeText( Text=text)

    def Save(self, Filename):
        self.obDoc.SaveAs(Filename)
    def Quit(self):
        self.obWord.Quit()

class QueryClientView(QtWidgets.QMainWindow,QPushButton):
    def __init__(self,data):
        super(QueryClientView, self).__init__()
        uic.loadUi("QueryClientView.ui", self).show()
        self.loadData(data)
    def loadData(self,data):
        self.setFixedWidth(511)
        self.setFixedWidth(601)
        self.ViewText.setText(data)

class QueryClientTable(QtWidgets.QMainWindow,QPushButton):
    def __init__(self,query,type):
        super(QueryClientTable, self).__init__()
        uic.loadUi("QueryClientTable.ui", self).show()
        self.loadtable(query,type)

    def loadtable(self,query,type):
        #this is done at server
        try:
            s.sendall(b'sqlQuery')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return
        stringServer=query+" "+type

        try:
            s.sendall(bytes(stringServer, "utf8"))
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return

        try:
            size = s.recv(1024)
            size = size.decode('utf-8')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return

        try:
            data = s.recv(int(size))
            data = data.decode('utf-8')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()


        #client receive string
        myTuple=data.split("|")
        number=myTuple[len(myTuple)-1]
        myTuple.remove(number)
        self.tableWidget.setRowCount(int(number))
        tableIndex = 0
        for row in myTuple:
            line=row.split("#")
            self.tableWidget.setItem(tableIndex, 0, QtWidgets.QTableWidgetItem(line[0]))
            self.tableWidget.setItem(tableIndex, 1, QtWidgets.QTableWidgetItem(line[1]))
            self.tableWidget.setItem(tableIndex, 2, QtWidgets.QTableWidgetItem(line[2]))
            self.tableWidget.setItem(tableIndex, 3, QtWidgets.QTableWidgetItem(line[3]))
            self.tableWidget.setItem(tableIndex, 4, QtWidgets.QTableWidgetItem(line[4]))
            tableIndex += 1

class QueryClient(QDialog):
    def __init__(self):
        super(QueryClient,self).__init__()
        loadUi("QueryClient.ui",self)
        self.SearchButton.clicked.connect(self.SearchFunction)
        self.ViewButton.clicked.connect(self.ViewFunction)
        self.DisconnectButton.clicked.connect(self.DisconnectFunction)
        self.DownLoadButton.clicked.connect(self.DownloadFunction)
    def DownloadFunction(self):
        global data
        type=str(self.comboBox.currentText())
        print(type)
        command = self.Command.text()
        comSplit = command.split(' ', 1)
        if (len(comSplit) != 0):
            if (comSplit[0] == "F_ID"):
                ID = comSplit[1]
                if ID.isdigit():
                    try:
                        s.sendall(b'Download')
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()
                        return
                    filename=ID+".txt"
                    try:
                        s.sendall(bytes(filename, "utf8"))
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()
                        return

                    try:
                        size = s.recv(1024)
                        size = size.decode('utf-8')
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()
                        return

                    try:
                        data = s.recv(int(size))
                        data = data.decode('utf-8')
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()

                    if(data=="Not Found"):
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Can't find file in library")
                        retval = msg.exec_()
                    else:
                        filename = 'DownloadFile//' + ID + type
                        if(type=='.txt'):
                            try:
                                file = open(filename, "w")
                                file.write(data)
                                file.close()
                            except:
                                msg = QtWidgets.QMessageBox()
                                msg.setIcon(QtWidgets.QMessageBox.Critical)
                                msg.setText("Download failed")
                                retval = msg.exec_()
                                return
                            msg = QtWidgets.QMessageBox()
                            msg.setText("Download Successfully !")
                            retval = msg.exec_()
                        elif(type=='.pdf'):
                            pdf = FPDF()
                            pdf.add_page()
                            pdf.set_font("Arial", size=15)
                            pdf.cell(200, 10, txt=data,ln=2, align='C')
                            pdf.output(filename)
                            msg = QtWidgets.QMessageBox()
                            msg.setText("Download Successfully !")
                            retval = msg.exec_()
                        elif (type=='.doc'):
                            markfile = 'DownloadFile\\MarkForDoc.txt'
                            f = open(markfile)
                            path = os.path.realpath(f.name)
                            print(path)
                            filepath = path.replace(markfile, '')
                            print(filepath)
                            obWord = CWordAutomate()
                            obWord.Write(data, "Courier New", 10)
                            obWord.Save(filepath + "\\DownloadFile\\"+ID+".doc")
                            obWord.Quit()
                            msg = QtWidgets.QMessageBox()
                            msg.setText("Download Successfully !")
                            retval = msg.exec_()
                        elif(type=='.docx'):
                            try:
                                myfile=docx.Document()
                                myfile.add_paragraph(data)
                                myfile.save(filename)
                            except:
                                msg = QtWidgets.QMessageBox()
                                msg.setIcon(QtWidgets.QMessageBox.Critical)
                                msg.setText("Download failed")
                                retval = msg.exec_()
                                return
                            msg = QtWidgets.QMessageBox()
                            msg.setText("Download Successfully !")
                            retval = msg.exec_()


    def DisconnectFunction(self):
        try:
           s.sendall(b'Disconnect')
           s.close()
           connect=Connect()
           widget.addWidget(connect)
           widget.setFixedWidth(347)
           widget.setFixedHeight(130)
           widget.setCurrentIndex(widget.currentIndex() + 1)
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()

    def ViewFunction(self):
        command = self.Command.text()
        comSplit = command.split(' ', 1)
        if (len(comSplit) != 0):
            if (comSplit[0] == "F_ID"):
                ID = comSplit[1]
                if ID.isdigit():
                    try:
                        s.sendall(b'View')
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()
                        return
                    filename=ID+".txt"
                    try:
                        s.sendall(bytes(filename, "utf8"))
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()
                        return

                    try:
                        size = s.recv(1024)
                        size = size.decode('utf-8')
                        print(size)
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()
                        return

                    try:
                        data = s.recv(int(size))
                        data = data.decode('utf-8')
                        print(data)
                    except:
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Connect failed!")
                        retval = msg.exec_()

                    if(data=="Not Found"):
                        msg = QtWidgets.QMessageBox()
                        msg.setIcon(QtWidgets.QMessageBox.Critical)
                        msg.setText("Can't find file in library")
                        retval = msg.exec_()
                    else:
                        self.view = QueryClientView(data)
                        self.view.show()
                else:
                    msg = QtWidgets.QMessageBox()
                    msg.setText("ID must be number")
                    retval = msg.exec_()
            else:
                msg = QtWidgets.QMessageBox()
                msg.setText("Wrong input for viewing")
                retval = msg.exec_()

    def SearchFunction(self):
        command=self.Command.text()
        comSplit=command.split(' ', 1)
        if(len(comSplit)!=0):
            if (comSplit[0] != "F_ID" and comSplit[0] != "F_Name" and comSplit[0] != "F_Author" and comSplit[0] != "F_Type") :
                msg = QtWidgets.QMessageBox()
                msg.setIcon(QtWidgets.QMessageBox.Critical)
                msg.setText("Wrong input")
                retval = msg.exec_()
                return False
            else:
                if(comSplit[0]=="F_ID"):
                    ID = comSplit[1]
                    if ID.isdigit():
                       self.IDfunct(ID)
                    else:
                        msg = QtWidgets.QMessageBox()
                        msg.setText("ID must be number")
                        retval = msg.exec_()
                else:
                    query=comSplit[1]
                    if(comSplit[0]=="F_Name"):
                        type="name"
                    elif(comSplit[0]=="F_Author"):
                        type="author"
                    else:
                        type = "type"
                    self.table=QueryClientTable(query,type)
                    self.table.show()
        else:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Wrong input")
            retval = msg.exec_()
            return False

    def IDfunct(self,ID):
        # send to sever here
        try:
            s.sendall(b'sqlQueryID')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return

        sqlcom = "SELECT * ,count (*) FROM book where ID=" + ID
        try:
            s.sendall(bytes(sqlcom, "utf8"))
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return

        try:
        # receive text here
            size = s.recv(1024)
            size = size.decode('utf-8')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return

        try:
            data = s.recv(int(size))
            data = data.decode('utf-8')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return

        text=data.split(" ",1)
        if (int(text[0])):
            self.DemoView.setText(text[1])
        else:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Can't find file in library")
            retval = msg.exec_()
            return 0


class Connect(QDialog):
    def __init__(self):
        super(Connect,self).__init__()
        loadUi("ConnectServer.ui",self)
        self.ConnectButton.clicked.connect(self.ConnectFunction)
    def ConnectFunction(self):
        #Input here
        HOST = self.InputIP.text()
        PORT = 8000
        if (HOST == "" ):
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("IP and Port must be filled!")
            retval = msg.exec_()
            return
        global s
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
     
        try:
            s.connect((HOST, PORT))
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed !")
            retval = msg.exec_()
            return
        #if pass

        queryClient = QueryClient()
        msg = QtWidgets.QMessageBox()
        msg.setText("Established connection successfully!")
        retval = msg.exec_()
        login=Login()
        widget.setFixedWidth(473)
        widget.setFixedHeight(360)
        widget.addWidget(login)
        widget.setCurrentIndex(widget.currentIndex() + 1)

#Create Login action
class Login(QDialog):
    def __init__(self):
        super(Login,self).__init__()
        loadUi("login.ui",self)
        self.loginbutton.clicked.connect(self.loginfunction)
        self.password.setEchoMode(QtWidgets.QLineEdit.Password)
        self.createaccbutton.clicked.connect(self.gotocreate)

    def loginfunction(self):
        try:
            s.sendall(b'login')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return
        user=self.user.text()
        password=self.password.text()
        tk = user + ' ' + password
        try:
            s.sendall(bytes(tk, "utf8"))
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Login failed !")
            retval = msg.exec_()
            return
        try:
            s.settimeout(5)
            data = s.recv(1024)
            data = data.decode('utf-8')
            s.settimeout(None)
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Login failed !")
            retval = msg.exec_()
            return
        Switcher(int(data))


    def gotocreate(self):
        createacc=CreateAcc()
        widget.setFixedWidth(480)
        widget.setFixedHeight(620)
        widget.addWidget(createacc)
        widget.setCurrentIndex(widget.currentIndex()+1)


#Create Create Account Action
class CreateAcc(QDialog):
    def __init__(self):
        super(CreateAcc,self).__init__()
        loadUi("createacc.ui",self)
        self.signupbutton.clicked.connect(self.createaccfunction)
        self.password.setEchoMode(QtWidgets.QLineEdit.Password)
        self.confirmpass.setEchoMode(QtWidgets.QLineEdit.Password)

    def createaccfunction(self):
        try:
            s.sendall(b'create')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Connect failed!")
            retval = msg.exec_()
            return
        user = self.user.text()
        password = self.password.text()
        confirm = self.confirmpass.text()
        tk = user + ' ' + password + ' ' + confirm
        s.sendall(bytes(tk, "utf8"))

        try:
            data = s.recv(1024)
            data = data.decode('utf-8')
        except:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Create account failed !")
            retval = msg.exec_()
            return

        if int(data) == 0:
            msg = QtWidgets.QMessageBox()
            msg.setIcon(QtWidgets.QMessageBox.Critical)
            msg.setText("Existed User")
            retval = msg.exec_()
        if int(data) == 1:
            msg = QtWidgets.QMessageBox()
            msg.setText("Registered Successfully !")
            retval = msg.exec_()
            login=Login()
            widget.addWidget(login)
            widget.setCurrentIndex(widget.currentIndex()+1)


#Switch to different dialog
def Switcher(i):
    if i==-1:
        msg = QtWidgets.QMessageBox()
        msg.setIcon(QtWidgets.QMessageBox.Critical)
        msg.setText("Invalid login")
        retval = msg.exec_()
        return False
    elif i==0:
        msg = QtWidgets.QMessageBox()
        msg.setIcon(QtWidgets.QMessageBox.Critical)
        msg.setText("Wrong Password")
        retval = msg.exec_()
        return False
    elif i==1:
        queryClient=QueryClient()
        msg = QtWidgets.QMessageBox()
        msg.setText("Login Success")
        retval = msg.exec_()
        widget.addWidget(queryClient)
        widget.setFixedWidth(660)
        widget.setFixedHeight(600)
        widget.setCurrentIndex(widget.currentIndex() + 1)


#main
if __name__ == "__main__":
    app=QApplication(sys.argv)
    mainwindow=Connect()
    widget=QtWidgets.QStackedWidget()
    widget.addWidget(mainwindow)
    widget.setFixedWidth(347)
    widget.setFixedHeight(130)
    widget.show()
    app.exec_()


